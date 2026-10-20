#!/usr/bin/env python3
"""Generate SDT 202 Team 4 initial presentation slides (week 8 scope)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "SDT202-project-team4-initial-slides.GENERATED.pptx"
PROTECTED = ROOT / "SDT202-project-team4-initial-slides.pptx"

# Palette: academic slate + teal accent (avoid purple / cream templates)
NAVY = RGBColor(0x1B, 0x2A, 0x41)
TEAL = RGBColor(0x1D, 0x6F, 0x70)
INK = RGBColor(0x24, 0x28, 0x2E)
MUTED = RGBColor(0x5C, 0x65, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
RULE = RGBColor(0xD0, 0xD7, 0xDE)


def rgb(shape_fill, color):
    shape_fill.solid()
    shape_fill.fore_color.rgb = color


def set_run(run, size=18, bold=False, color=INK, italic=False, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name


def add_textframe(shape, paragraphs, *, default_size=18, default_color=INK, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(paragraphs):
        if isinstance(item, str):
            text, size, bold, color, italic = item, default_size, False, default_color, False
        else:
            text = item.get("text", "")
            size = item.get("size", default_size)
            bold = item.get("bold", False)
            color = item.get("color", default_color)
            italic = item.get("italic", False)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = text
        set_run(run, size=size, bold=bold, color=color, italic=italic)


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    shape.line.fill.background()
    rgb(shape.fill, color)
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_top_bar(slide):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
    )
    bar.line.fill.background()
    rgb(bar.fill, TEAL)


def add_footer(slide, page, total):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(10), Inches(0.35))
    add_textframe(
        box,
        [
            {
                "text": "SDT 202 · Team 4 · Initial presentation · [Public repository URL]",
                "size": 11,
                "color": MUTED,
            }
        ],
    )
    num = slide.shapes.add_textbox(Inches(11.6), Inches(7.05), Inches(1.2), Inches(0.35))
    add_textframe(
        num,
        [{"text": f"{page}/{total}", "size": 11, "color": MUTED}],
        align=PP_ALIGN.RIGHT,
    )


def title_block(slide, title, subtitle=None):
    add_top_bar(slide)
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.45), Inches(0.12), Inches(0.55)
    )
    accent.line.fill.background()
    rgb(accent.fill, TEAL)
    t = slide.shapes.add_textbox(Inches(0.9), Inches(0.4), Inches(11.5), Inches(0.7))
    add_textframe(t, [{"text": title, "size": 28, "bold": True, "color": NAVY}])
    if subtitle:
        s = slide.shapes.add_textbox(Inches(0.9), Inches(1.0), Inches(11.5), Inches(0.4))
        add_textframe(s, [{"text": subtitle, "size": 14, "color": MUTED}])


def bullets(slide, items, left=0.9, top=1.5, width=11.5, height=5.2, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, str):
            text, lvl, bold = item, 0, False
        else:
            text = item.get("text", "")
            lvl = item.get("level", 0)
            bold = item.get("bold", False)
        prefix = "    ◦ " if lvl else "• "
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.clear()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = prefix + text
        set_run(run, size=size - (2 if lvl else 0), bold=bold, color=INK)


def code_box(slide, lines, left=0.9, top=1.55, width=11.5, height=4.8):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    rgb(shape.fill, LIGHT)
    shape.line.color.rgb = RULE
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_top = Inches(0.2)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        set_run(run, size=14, color=INK, name="Consolas")


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slides_spec = []

    # Collect builders then stamp footers with total
    def slide(fn):
        slides_spec.append(fn)

    @slide
    def s_title(sl, page, total):
        add_bg(sl, NAVY)
        bar = sl.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.9), Inches(13.333), Inches(0.6)
        )
        bar.line.fill.background()
        rgb(bar.fill, TEAL)
        box = sl.shapes.add_textbox(Inches(0.9), Inches(1.8), Inches(11.5), Inches(3.5))
        add_textframe(
            box,
            [
                {"text": "SDT 202", "size": 18, "color": TEAL, "bold": True},
                {
                    "text": "Gradient Descent Algorithm",
                    "size": 36,
                    "bold": True,
                    "color": WHITE,
                },
                {
                    "text": "RetentionRadar — initial presentation (≈30% progress)",
                    "size": 20,
                    "color": WHITE,
                },
                {"text": " ", "size": 12, "color": WHITE},
                {"text": "Team 4", "size": 18, "color": WHITE, "bold": True},
                {"text": "[Student A Full Name]  ·  [Student B Full Name]", "size": 16, "color": WHITE},
                {"text": "Group [Group Code]  ·  American University Kyiv", "size": 14, "color": WHITE},
                {"text": "Instructor: [Instructor Full Name]", "size": 14, "color": WHITE},
            ],
        )
        foot = sl.shapes.add_textbox(Inches(0.9), Inches(7.0), Inches(11), Inches(0.35))
        add_textframe(
            foot,
            [{"text": f"[Public repository URL]   ·   {page}/{total}", "size": 12, "color": WHITE}],
        )

    @slide
    def s_agenda(sl, page, total):
        add_bg(sl)
        title_block(sl, "Agenda", "Initial presentation scope from the course brief")
        bullets(
            sl,
            [
                "1. Problem description",
                "2. Classical / general algorithm",
                "3. Pseudocode",
                "4. Real-world situation (RetentionRadar)",
                "5. Progress, repository, risks, next steps",
                {"text": "Each member presents a part — cameras on if online", "bold": True},
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_roles(sl, page, total):
        add_bg(sl)
        title_block(sl, "Who presents what", "Fill in names before rehearsal")
        bullets(
            sl,
            [
                "[Student A]: problem + classical algorithm intuition + convergence idea",
                "[Student B]: pseudocode walk-through + real-world application",
                "Together: progress status, repository link, risks, Q&A",
                {"text": "Brief reminder: every member must be able to answer on any part", "bold": True},
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_problem(sl, page, total):
        add_bg(sl)
        title_block(sl, "1. Problem", "Unconstrained smooth minimization")
        bullets(
            sl,
            [
                "Given a differentiable objective f : ℝᵈ → ℝ, find x that minimizes f(x)",
                "Setting ∇f(x) = 0 often yields a nonlinear system with no closed form",
                "Need a numerical method that uses only first-order information",
                "Gradient descent: repeatedly step opposite the gradient",
                {"text": "Negative gradient = steepest local descent direction (Euclidean norm)", "bold": True},
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_classical(sl, page, total):
        add_bg(sl)
        title_block(sl, "2. Classical algorithm", "Batch gradient descent")
        bullets(
            sl,
            [
                "Update rule:",
                {"text": "      xₖ₊₁  =  xₖ  −  α · ∇f(xₖ)", "bold": True},
                "α > 0 is a fixed step size (learning rate)",
                "“Batch” = each gradient uses the full objective / full dataset",
                "Stop when ‖∇f(x)‖ is small, or after a maximum number of updates",
                "If α is too large, iterates can diverge; if too small, progress is slow",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_theory(sl, page, total):
        add_bg(sl)
        title_block(sl, "Classical algorithm — why step size matters", "L-smooth convex case (standard lecture bound)")
        bullets(
            sl,
            [
                "If ∇f is L-Lipschitz (L-smooth) and f is convex with a minimizer x*",
                "Choose 0 < α ≤ 1/L",
                {"text": "Then:  f(xₖ) − f(x*)  ≤  ‖x₀ − x*‖² / (2 α k)   →   O(1/k)", "bold": True},
                "Our program stops on gradient norm (not the same as objective-gap ε)",
                "Complexity per iteration depends on the cost of ∇f (for logistic: Θ(m·d))",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_pseudo1(sl, page, total):
        add_bg(sl)
        title_block(sl, "3. Pseudocode", "Matches GradientDescent.minimize in our Java code")
        code_box(
            sl,
            [
                "GRADIENT-DESCENT(f, x0, alpha, tolerance, maxUpdates)",
                "    x ← copy(x0)",
                "    updates ← 0",
                "    repeat",
                "        value ← f(x)",
                "        if nonfinite(x/value) or magnitude-guard fails",
                "            return DIVERGED",
                "        g ← ∇f(x)",
                "        if nonfinite(g) or nonfinite(‖g‖)",
                "            return DIVERGED",
                "        if ‖g‖ ≤ tolerance",
                "            return x, CONVERGED, updates",
                "        if updates = maxUpdates",
                "            return x, MAX_ITERATIONS, updates",
                "        x ← x − alpha · g",
                "        updates ← updates + 1",
            ],
            height=5.1,
        )
        add_footer(sl, page, total)

    @slide
    def s_pseudo2(sl, page, total):
        add_bg(sl)
        title_block(sl, "Pseudocode — notes", "Implementation semantics")
        bullets(
            sl,
            [
                "Stopping test is before an update → optimal start needs 0 updates",
                "Also check the point produced by the final allowed update",
                "All coordinates share the same gradient vector g (true batch GD)",
                "Magnitude guard (practical): value > 10⁶ · (|f(x₀)| + 1) → DIVERGED",
                "Java API: minimize(ObjectiveFunction, double[] x0, OptimizerConfig)",
                "Driver RetentionRadarApp calls GradientDescent.minimize(...)",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_example(sl, page, total):
        add_bg(sl)
        title_block(sl, "Worked example", "f(x) = x²,  ∇f = 2x,  L = 2,  x₀ = 5")
        bullets(
            sl,
            [
                "α = 0.1  →  5, 4, 3.2, 2.56, …  (contracts by 0.8 each step)",
                "α = 0.9  →  oscillates while contracting",
                "α = 1.1  →  magnitude grows (step too large vs 2/L)",
                "Demo in code: quadratic bowl + ASCII contours (menu option 8)",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_realworld(sl, page, total):
        add_bg(sl)
        title_block(sl, "4. Real-world situation", "RetentionRadar — educational early-warning demo")
        bullets(
            sl,
            [
                "Problem: estimate risk of eventual Dropout vs Graduate",
                "Using information available by the end of semester one",
                "Fit logistic regression by minimizing regularized logistic loss with GD",
                "Practical value: rank held-out students so advisors know whom to review first",
                {
                    "text": "Real data: UCI “Predict Students’ Dropout and Academic Success” "
                    "(dataset #697, Realinho et al.; CC BY 4.0) [3][4]",
                    "bold": True,
                },
                "Our extract: 7 first-semester features; 4,424 rows; Enrolled dropped at load",
                "Also: seeded synthetic cohort for offline demos without the CSV",
                {"text": "Limits: educational prototype — not AUK operational data or causal proof", "bold": True},
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_data_cite(sl, page, total):
        add_bg(sl)
        title_block(sl, "Data source (citation)", "What UCI #697 is, in one slide")
        bullets(
            sl,
            [
                "Public dataset from the UCI Machine Learning Repository",
                "Full name: Predict Students’ Dropout and Academic Success",
                "Authors: Realinho, Vieira Martins, Machado, Baptista (and related Data 2022 paper)",
                "Labels: Dropout / Graduate / Enrolled (we keep Dropout vs Graduate only)",
                "Licence: Creative Commons Attribution 4.0 (CC BY 4.0) — attribute + note changes",
                "Project file: data/students-first-semester.csv (+ DATA-SOURCE.md attribution)",
                "URL: archive.ics.uci.edu/dataset/697/…  (see References slide)",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_pipeline(sl, page, total):
        add_bg(sl)
        title_block(sl, "How the Java program uses the algorithm", "Main program must call the algorithm function")
        bullets(
            sl,
            [
                "LogisticObjective implements ObjectiveFunction (value + gradient)",
                "RetentionRadarApp prepares data → calls GradientDescent.minimize",
                "Then scores held-out records, metrics, and optional ranking",
                "Packages: optimization / model / data / app / verify",
                "Stack: plain Java 17, IntelliJ, no external libraries",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_progress(sl, page, total):
        add_bg(sl)
        title_block(sl, "Progress for week 8 (initial)", "Brief asks for ≈30% — honest status")
        bullets(
            sl,
            [
                "Done / in progress toward the initial deliverable:",
                {"text": "Classical GD implemented as a reusable Java function", "level": 1},
                {"text": "Real-world logistic application wired to call that function", "level": 1},
                {"text": "Public GitHub repository + IntelliJ project structure", "level": 1},
                {"text": "Report draft + these slides for Canvas (24h before talk)", "level": 1},
                "Still for later / final (not the focus of this talk):",
                {"text": "Improved algorithm (Nesterov), full screenshots, comments/header polish", "level": 1},
                {"text": "Final report parts, abstract, effort analysis, live demo polish", "level": 1},
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_risks(sl, page, total):
        add_bg(sl)
        title_block(sl, "Project risks", "Called out for the initial presentation")
        bullets(
            sl,
            [
                "Step size too large → non-finite loss / divergence (mitigate with safe α = 1/B)",
                "Ill-conditioned features → many iterations (standardize using training only)",
                "Feature timing misunderstanding → invalid “early warning” claims",
                "Single split oversold as deployment performance",
                "Repo visibility / clickable link / contribution evidence before final",
                "Schedule pressure — start early; rehearse with both members speaking",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_repo(sl, page, total):
        add_bg(sl)
        title_block(sl, "Repository & submission checklist", "Initial Canvas packet")
        bullets(
            sl,
            [
                "One GitHub repository (local + remote): [Public repository URL]",
                "Make the report link clickable and verify it opens logged out",
                "Submit 24 hours before the presentation:",
                {"text": "Word report: SDT202-project-team4.docx (template parts for initial)", "level": 1},
                {"text": "These slides", "level": 1},
                "Initial report parts: title · introduction · classical algorithm · pseudocode · real-world situation",
                "Entire team present; each member presents their part",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_next(sl, page, total):
        add_bg(sl)
        title_block(sl, "Next steps (toward final)", "Not required content for today’s deep dive")
        bullets(
            sl,
            [
                "Human-authored driver header + necessary comments (no AI comments)",
                "Describe and compare the improved algorithm (Nesterov)",
                "Screenshots titled like “Part 2.1. Screenshots: …” for every coding part",
                "Complete report sections, abstract, effort analysis from real commits",
                "Final presentation ≈ week 15: live demo, every member speaks",
            ],
        )
        add_footer(sl, page, total)

    @slide
    def s_refs(sl, page, total):
        add_bg(sl)
        title_block(sl, "References", "Cited in this presentation")
        bullets(
            sl,
            [
                "[1] Tibshirani R. Gradient Descent (CMU convex opt. notes).",
                "[2] Su, Boyd, Candes. ODE analysis of Nesterov’s accelerated gradient.",
                "[3] Realinho V. et al. Predict Students’ Dropout and Academic Success. "
                "UCI ML Repository, dataset #697. "
                "https://archive.ics.uci.edu/dataset/697/predict+students+dropout+and+academic+success",
                "[4] Realinho et al. Predicting Student Dropout and Academic Success. "
                "Data 2022, 7(11), 146. https://www.mdpi.com/2306-5729/7/11/146",
                "[5] Shalizi C. Logistic Regression (CMU lecture notes).",
                "[6] Creative Commons Attribution 4.0 (CC BY 4.0).",
            ],
            size=15,
        )
        add_footer(sl, page, total)

    @slide
    def s_thanks(sl, page, total):
        add_bg(sl, NAVY)
        box = sl.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(2.5))
        add_textframe(
            box,
            [
                {"text": "Thank you", "size": 40, "bold": True, "color": WHITE},
                {"text": "Questions?", "size": 24, "color": WHITE},
                {"text": " ", "size": 12, "color": WHITE},
                {
                    "text": "Team 4 · [Student A] · [Student B] · [Public repository URL]",
                    "size": 16,
                    "color": WHITE,
                },
            ],
            align=PP_ALIGN.CENTER,
        )
        # center-ish by wide box
        for shape in sl.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER

    total = len(slides_spec)
    for i, fn in enumerate(slides_spec, start=1):
        sl = blank_slide(prs)
        fn(sl, i, total)

    prs.save(OUT)
    print(f"Wrote {OUT} ({total} slides)")
    if PROTECTED.exists():
        print(
            f"Note: left hand-edited {PROTECTED.name} untouched. "
            f"Copy from {OUT.name} only if you intend to replace it."
        )


if __name__ == "__main__":
    build()
