#!/usr/bin/env python3
"""Generate the SDT 202 Team 4 project report (Word)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "SDT202-project-team4-report.docx"

INSTRUCTOR = "Dr. Valeriy Kolesnikov"
STUDENT_A = "Uliana Matviienko"
STUDENT_B = "Andrii Moroz"
GROUP = "BSEAI 2025"
TEAM = "Team 4"
# Clickable in the report body (not on the title page). Official brief requires a
# clickable repository link somewhere in the report; the Team1 example title page
# did not list a repository.
REPO_URL = "https://github.com/shalltearfrost/RetentionRadar"


def set_run_font(run, *, size=12, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if color is not None:
        run.font.color.rgb = color


def style_paragraph(p, *, align="justify", space_after=8, space_before=0, line=1.15):
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line


def add_para(doc, text="", *, bold=False, size=12, align="justify", space_after=8, italic=False):
    p = doc.add_paragraph()
    style_paragraph(p, align=align, space_after=space_after)
    if text:
        set_run_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def configure_heading_styles(doc):
    for style_name, size, center in (("Heading 1", 14, True), ("Heading 2", 12, False)):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(10)
        if center:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")


def add_h1(doc, text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        set_run_font(run, size=14, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_h2(doc, text):
    p = doc.add_heading(text, level=2)
    for run in p.runs:
        set_run_font(run, size=12, bold=True)
    return p


def add_figure_placeholder(doc, caption):
    add_para(
        doc,
        "[SCREENSHOT PLACEHOLDER — insert image here]",
        align="center",
        italic=True,
        size=11,
        space_after=4,
    )
    add_para(doc, caption, align="center", bold=True, size=11, space_after=12)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        style_paragraph(p, align="justify", space_after=4)
        set_run_font(p.add_run(item), size=12)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_run_font(cell.paragraphs[0].add_run(h), bold=True, size=10)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            set_run_font(cell.paragraphs[0].add_run(str(val)), size=10)
    doc.add_paragraph()


def add_hyperlink(paragraph, text, url):
    """Add a clickable hyperlink run to an existing paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")  # 12 pt
    rPr.append(sz)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rPr.append(rFonts)
    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=11)
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r.append(fld_begin)

    run2 = paragraph.add_run()
    set_run_font(run2, size=11)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run2._r.append(instr)

    run3 = paragraph.add_run()
    set_run_font(run3, size=11)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run3._r.append(fld_end)


def setup_page_numbers(doc):
    """Automatic PAGE field in footer; title page has no number."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    # Empty first-page footer (title page)
    first = section.first_page_footer
    first.is_linked_to_previous = False
    if first.paragraphs:
        first.paragraphs[0].clear()

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number_field(p)


def add_toc_field(doc):
    """Insert a Word TOC field (hyperlinked). Updates when the document is opened."""
    p = doc.add_paragraph()
    style_paragraph(p, align="left", space_after=6)
    run = p.add_run()
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r.append(fld_begin)

    run2 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    # \o "1-2" = heading levels; \h = hyperlinks; \z = hide tab in web view; \u = outline levels
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    run2._r.append(instr)

    run3 = p.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(fld_sep)

    run4 = p.add_run(
        "(Table of contents — in Word/LibreOffice: right-click → Update Field / Update Index)"
    )
    set_run_font(run4, size=11, italic=True, color=RGBColor(0x66, 0x66, 0x66))

    run5 = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run5._r.append(fld_end)


def enable_update_fields_on_open(doc):
    settings = doc.settings.element
    # Remove existing updateFields if present
    for child in list(settings):
        if child.tag == qn("w:updateFields"):
            settings.remove(child)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_label_value_line(doc, label, value):
    p = doc.add_paragraph()
    style_paragraph(p, align="left", space_after=4)
    set_run_font(p.add_run(f"{label}"), size=12, bold=True)
    set_run_font(p.add_run(f"\t{value}"), size=12)


def build():
    doc = Document()
    configure_heading_styles(doc)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)

    setup_page_numbers(doc)
    enable_update_fields_on_open(doc)

    # ----- Title page (no repository — matches Team1 example layout) -----
    for _ in range(2):
        add_para(doc, "", space_after=0)
    add_para(doc, "American University Kyiv", align="center", bold=True, size=14, space_after=6)
    add_para(
        doc,
        "SDT 202: Design and Analysis of Data Structures and Algorithms",
        align="center",
        bold=True,
        size=13,
        space_after=48,
    )
    add_para(
        doc,
        "Project report",
        align="center",
        bold=True,
        size=14,
        space_after=6,
    )
    add_para(
        doc,
        "“Gradient Descent Algorithm — RetentionRadar”",
        align="center",
        bold=True,
        size=14,
        space_after=48,
    )

    add_label_value_line(doc, "Instructor:", INSTRUCTOR)
    add_label_value_line(doc, "Students:", STUDENT_A)
    p = doc.add_paragraph()
    style_paragraph(p, align="left", space_after=4)
    set_run_font(p.add_run(f"\t{STUDENT_B}"), size=12)
    add_label_value_line(doc, "Group:", GROUP)
    add_label_value_line(doc, "Team:", TEAM)

    for _ in range(6):
        add_para(doc, "", space_after=0)
    add_para(doc, "Kyiv 2026", align="center", bold=True, size=12, space_after=0)
    doc.add_page_break()

    # ----- Contents (auto TOC field) -----
    add_h1(doc, "CONTENTS")
    add_para(
        doc,
        "The entries below are a Word/LibreOffice table of contents field. "
        "They link to headings and show page numbers after the field is updated.",
        align="justify",
        italic=True,
        size=11,
    )
    add_toc_field(doc)
    doc.add_page_break()

    # ----- 1 Introduction -----
    add_h1(doc, "1 INTRODUCTION")
    add_para(
        doc,
        "The aim of this work is to implement classical batch gradient descent and an "
        "accelerated improvement (Nesterov’s method) in plain Java 17, and to apply them "
        "to a concrete smooth convex objective: binary logistic regression for an educational "
        "student-retention demonstration called RetentionRadar.",
    )
    add_para(
        doc,
        "Gradient descent is a first-order iterative method for minimizing a differentiable "
        "objective f : R^d → R. When a closed-form solution of ∇f(x) = 0 is unavailable, "
        "repeated updates of the form x ← x − α∇f(x) provide a practical numerical procedure. "
        "Under standard smoothness and convexity assumptions, classical gradient descent "
        "admits an O(1/k) objective-gap bound; Nesterov acceleration improves the convex "
        "rate to O(1/k²) for ordinary iterates under an appropriate step-size choice.",
    )
    add_para(
        doc,
        "The application minimizes regularized logistic loss on synthetic data and on an "
        "attributed extract of the UCI Machine Learning Repository dataset "
        "“Predict Students’ Dropout and Academic Success” (dataset #697) by Realinho et al. [3]. "
        "The related open-access descriptor is Realinho et al. (2022) [4]. "
        "The binary task is eventual Dropout versus Graduate using seven "
        "first-semester features. The program trains models, compares optimizers, evaluates "
        "held-out metrics, ranks records, supports an illustrative cost-based decision "
        "threshold, and draws ASCII objective contours with sampled optimizer paths. "
        "Dataset attribution, licence (CC BY 4.0) [6] and the exact column changes for this "
        "project are recorded in data/DATA-SOURCE.md.",
    )
    p = doc.add_paragraph()
    style_paragraph(p)
    set_run_font(
        p.add_run(
            "Collective work uses a single GitHub repository [7]. Project repository: "
        ),
        size=12,
    )
    add_hyperlink(p, REPO_URL, REPO_URL)
    set_run_font(
        p.add_run(
            ". The implementation uses only the JDK (no Maven/Gradle and no external Java "
            "libraries). Interactive use is through a console menu; demonstration modes are "
            "also available via command-line flags."
        ),
        size=12,
    )
    doc.add_page_break()

    # ----- 2 Requirements -----
    add_h1(doc, "2 REQUIREMENTS")
    add_para(
        doc,
        "The primary requirement is to implement the gradient descent algorithm in Java "
        "and demonstrate it on a meaningful optimization problem, including an improved "
        "variant with comparison against the classical method.",
    )
    add_para(doc, "Project requirements addressed by this deliverable include:")
    add_bullets(
        doc,
        [
            "implementation of classical batch gradient descent with a shared optimizer interface;",
            "implementation of Nesterov accelerated gradient descent as the improvement;",
            "a reusable ObjectiveFunction abstraction (value and gradient);",
            "a stable logistic-regression objective with unregularized bias;",
            "synthetic cohort generation with a fixed seed for offline demos;",
            "loading of an attributed real CSV extract with strict schema checks;",
            "stratified train/test split and train-only feature standardization;",
            "held-out evaluation (accuracy, precision, recall, F1, AUC, log-loss, Brier, confusion);",
            "comparison of updates, objective values, gradient norms and evaluation call counts;",
            "ASCII contour visualization of optimizer trajectories;",
            "top-25 held-out ranking and an illustrative error-cost threshold control;",
            "dependency-free verification suite (58 checks);",
            "English interactive console menu and non-interactive demonstration modes;",
            "reproducibility with JDK 17+ and IntelliJ project descriptors.",
        ],
    )
    doc.add_page_break()

    # ----- 3 Design -----
    add_h1(doc, "3 DESIGN")
    add_h2(doc, "3.1 Classical gradient descent")
    add_para(
        doc,
        "The vector update is x_{k+1} = x_k − α ∇f(x_k). For an L-smooth convex objective "
        "with 0 < α ≤ 1/L and an attained minimizer, the standard bound is "
        "f(x_k) − f(x*) ≤ ‖x_0 − x*‖² / (2 α k), i.e. an O(1/k) objective-gap rate. "
        "The implementation stops on gradient-norm tolerance (a different criterion from "
        "an objective-gap benchmark).",
    )
    add_para(doc, "Pseudocode matching GradientDescent.minimize:")
    add_para(
        doc,
        "GRADIENT-DESCENT(f, x0, α, tolerance, maxUpdates)\n"
        "    x ← copy(x0); updates ← 0\n"
        "    repeat\n"
        "        value ← f(x); g ← ∇f(x)\n"
        "        if non-finite or magnitude-guard failure → return DIVERGED\n"
        "        if ‖g‖ ≤ tolerance → return x, CONVERGED, updates\n"
        "        if updates = maxUpdates → return x, MAX_ITERATIONS, updates\n"
        "        x ← x − α g; updates ← updates + 1",
        size=10,
        align="left",
    )
    add_para(
        doc,
        "Worked example: for f(x) = x², ∇f = 2x, L = 2, start x0 = 5. With α = 0.1 the "
        "iterates contract (5, 4, 3.2, …); with α = 1.1 the magnitude grows.",
    )

    add_h2(doc, "3.2 Improvement: Nesterov acceleration")
    add_para(
        doc,
        "NesterovGradientDescent evaluates the update gradient at a look-ahead point y and "
        "uses momentum β_k = (k−1)/(k+2) with 0 < α ≤ 1/L. Stopping, history and the "
        "returned point refer to the ordinary iterate x. For smooth convex objectives the "
        "ordinary-iterate gap bound is O(1/k²): "
        "f(x_k) − f(x*) ≤ 2‖x_0−x*‖² / (α (k+1)²).",
    )

    add_h2(doc, "3.3 Application objective and data")
    add_para(
        doc,
        "Logistic loss with softplus-stable evaluation, unregularized bias and slope L2 "
        "penalty λ is minimized. A conservative smoothness bound "
        "B = Σ‖x_i‖²/(4m) + λ yields α = 1/B. "
        "Real data come from the UCI dataset “Predict Students’ Dropout and Academic Success” "
        "[3], [4]: the bundled extract keeps 4 424 source rows and seven first-semester "
        "features; at load time 794 Enrolled outcomes are omitted and 3 630 "
        "Dropout/Graduate records are retained; a stratified 80/20 split (seed 202604) "
        "is used for evaluation.",
    )

    add_h2(doc, "3.4 Program architecture and interface")
    add_para(
        doc,
        "The program is a console application written in plain Java 17. There is no "
        "graphical UI toolkit; visualization is provided by ASCII contour plots printed "
        "to the terminal. Numeric output uses Locale.ROOT so decimal points remain "
        "stable across JVM locales.",
    )
    add_para(
        doc,
        "Source packages under auk.sdt202.team4 separate concerns: optimization "
        "(classical GD, Nesterov, configuration and results), model (logistic objective), "
        "data (CSV loader, synthetic cohort, split, standardizer), app (driver, evaluation, "
        "ranking, costs, contours, comparison), and verify (self-checks).",
    )
    add_para(
        doc,
        "Both optimizers expose the same entry point "
        "minimize(ObjectiveFunction, double[], OptimizerConfig) and return an "
        "OptimizationResult with the ordinary iterate, objective value, gradient norm, "
        "update count, stop reason and optional sampled history. Stopping uses the "
        "gradient-norm tolerance before each update and again after the final allowed update.",
    )
    add_para(
        doc,
        "The main interactive window of the program is the console menu shown after launch "
        "with no arguments. Replace the placeholder below with a screenshot of the menu "
        "at project root working directory.",
    )
    add_figure_placeholder(doc, "Figure 3.1 – Main console menu of RetentionRadarApp")

    add_para(
        doc,
        "Selecting option 1 loads the bundled UCI extract. The loader reports source rows, "
        "omitted Enrolled outcomes and retained Dropout/Graduate records.",
    )
    add_figure_placeholder(doc, "Figure 3.2 – Loading the bundled UCI cohort (menu option 1)")

    add_para(
        doc,
        "Options 3 and 4 train classical gradient descent and Nesterov acceleration on the "
        "current cohort. Option 5 trains both methods, prints a comparison table, held-out "
        "metrics and projected ASCII contours for two coefficient coordinates.",
    )
    add_figure_placeholder(doc, "Figure 3.3 – Classical gradient descent training summary (menu option 3)")
    add_figure_placeholder(doc, "Figure 3.4 – Nesterov training summary (menu option 4)")
    add_figure_placeholder(
        doc,
        "Figure 3.5 – Side-by-side comparison, held-out metrics and ASCII contours (menu option 5)",
    )

    add_para(
        doc,
        "Option 6 prints the top-25 held-out ranking by predicted risk score. Option 7 "
        "changes illustrative missed-positive and unnecessary-outreach costs; the decision "
        "threshold updates without retraining model coefficients.",
    )
    add_figure_placeholder(doc, "Figure 3.6 – Held-out ranking, top 25 (menu option 6)")
    add_figure_placeholder(doc, "Figure 3.7 – Cost inputs and updated threshold (menu option 7)")

    add_para(
        doc,
        "Option 8 runs the known quadratic comparison used to illustrate acceleration on a "
        "smooth strongly curved bowl. Option 9 runs the full verification suite.",
    )
    add_figure_placeholder(doc, "Figure 3.8 – Quadratic comparison and ASCII contours (menu option 8)")
    add_figure_placeholder(doc, "Figure 3.9 – Verification checks passing (menu option 9 / --verify)")

    add_para(
        doc,
        "IntelliJ project layout and the remote repository provide structural evidence of "
        "the deliverable organization.",
    )
    add_figure_placeholder(doc, "Figure 3.10 – IntelliJ project structure (src packages and data/)")
    add_figure_placeholder(doc, "Figure 3.11 – GitHub repository page showing contributors and project files")
    doc.add_page_break()

    # ----- 4 Diagrams -----
    add_h1(doc, "4 DIAGRAMS")
    add_para(
        doc,
        "Figure 4.1 summarizes interactive use-cases of the console application. "
        "Figure 4.2 summarizes the main classes and their relationships. Insert diagrams "
        "created in any UML tool; the captions below are reserved for those figures.",
    )
    add_figure_placeholder(doc, "Figure 4.1 – Use-case diagram")
    add_para(doc, "Suggested actors and use cases:")
    add_bullets(
        doc,
        [
            "Actor: Student / demonstrator.",
            "Load UCI cohort; Load synthetic cohort;",
            "Train classical GD; Train Nesterov;",
            "Compare methods and evaluate held-out data;",
            "Show ranking; Adjust decision costs;",
            "Run quadratic demo; Run verification checks;",
            "Exit application.",
        ],
    )
    add_figure_placeholder(doc, "Figure 4.2 – Class diagram")
    add_para(doc, "Suggested class groups for the diagram:")
    add_bullets(
        doc,
        [
            "optimization: ObjectiveFunction, OptimizerConfig, OptimizationResult, "
            "StopReason, GradientDescent, NesterovGradientDescent, VectorMath;",
            "model: LogisticObjective;",
            "data: SyntheticCohort, StudentCsvLoader, DataSplit, Standardizer;",
            "app: RetentionRadarApp, Evaluation, ComparisonReport, Watchlist, "
            "DecisionCosts, AsciiContour, QuadraticObjective;",
            "verify: SelfCheck, ModelCheck, AccelerationCheck, DataCheck, EvaluationCheck.",
        ],
    )
    doc.add_page_break()

    # ----- 5 Code -----
    add_h1(doc, "5 CODE WITH COMMENTS")
    add_para(
        doc,
        "Commented source code is provided in Appendix A. The assignment requires "
        "human-authored comments rather than generated line-by-line commentary. "
        "Team members must review every class and insert the required header "
        "and comments in their own words before final submission.",
    )
    add_para(
        doc,
        "The driver class is auk.sdt202.team4.app.RetentionRadarApp. Classical and "
        "accelerated optimizers live in the optimization package and share the "
        "ObjectiveFunction interface implemented by LogisticObjective and QuadraticObjective.",
    )
    doc.add_page_break()

    # ----- 6 Testing -----
    add_h1(doc, "6 TESTING REPORT")
    add_para(
        doc,
        "Functional behaviour was checked with the built-in verification mode "
        "(RetentionRadarApp --verify), which currently executes 58 dependency-free checks, "
        "and with recorded demonstration runs on synthetic data, the UCI extract and a "
        "known quadratic. Table 6.1 summarizes representative tests.",
    )
    add_para(doc, "Table 6.1 – Program tests", bold=True, align="center", size=11)
    add_table(
        doc,
        ["ID", "Description", "Steps", "Input", "Expected", "Actual", "0/1"],
        [
            [
                "V-01",
                "All verification checks pass",
                "Run --verify from project root",
                "—",
                "58 PASS lines; exit 0",
                "See Fig. 6.1",
                "1",
            ],
            [
                "V-02",
                "Classical GD reaches known quadratic minimum",
                "Run --quadratic or menu 8",
                "f=0.5(x²+100y²), start (1,1)",
                "CONVERGED near (0,0)",
                "See Fig. 6.2",
                "1",
            ],
            [
                "V-03",
                "Nesterov uses fewer updates than GD on ill-conditioned bowl",
                "Run --compare",
                "f=0.5(x²+1000y²), α=0.001",
                "Nesterov updates < GD updates",
                "GD 13809 / Nest. 3302 (recorded)",
                "1",
            ],
            [
                "V-04",
                "UCI loader retains resolved outcomes only",
                "Run --real or menu 1 then 5",
                "bundled CSV",
                "3630 retained; 794 Enrolled omitted",
                "Matches recorded run",
                "1",
            ],
            [
                "V-05",
                "Held-out metrics printed for both methods",
                "Menu 5 after UCI load",
                "seed 202604 split",
                "Accuracy/AUC etc. finite; both methods agree",
                "Acc. 0.88430, AUC 0.91960",
                "1",
            ],
            [
                "V-06",
                "Cost change updates threshold without retraining",
                "Menu 7 after compare",
                "missed=9, outreach=1",
                "Threshold falls; more eligible rows",
                "0.5→0.1; 232→558 eligible",
                "1",
            ],
            [
                "V-07",
                "Ranking lists top 25 held-out scores",
                "Menu 6",
                "fitted model present",
                "25 rows with scores/labels",
                "See Fig. 6.6",
                "1",
            ],
            [
                "V-08",
                "Invalid CSV schema rejected",
                "Load malformed file via --real path",
                "wrong headers",
                "Clear error; no crash",
                "Covered by DataCheck",
                "1",
            ],
            [
                "V-09",
                "Too-large step reports divergence",
                "Self-check magnitude/step cases",
                "unsafe α in checks",
                "StopReason DIVERGED or guarded",
                "Covered by SelfCheck",
                "1",
            ],
            [
                "V-10",
                "Fresh clone builds and verifies",
                "Clone repo; compile; --verify; --real",
                "JDK 17+",
                "Same counts/metrics as validation notes",
                "Documented in validation.md",
                "1",
            ],
        ],
    )

    add_para(
        doc,
        "Recorded UCI configuration: stratified 80/20 split (train 2904 / test 726, seed 202604), "
        "λ = 0.01, B = 2.01, α = 1/B, tolerance 1e−6, maximum 50 000 updates. "
        "Classical GD used 865 updates (866 gradient calls); Nesterov used 375 updates "
        "(749 gradient calls). Held-out confusion at threshold 0.5: TP=216, FP=16, TN=426, FN=68.",
    )
    add_para(
        doc,
        "Update-count reductions are not wall-clock speedups: Nesterov also evaluates "
        "look-ahead points. Results are for one fixed selected-cohort split and do not "
        "establish institutional deployment performance.",
    )

    add_figure_placeholder(doc, "Figure 6.1 – Verification output (--verify), all checks PASS")
    add_figure_placeholder(doc, "Figure 6.2 – Quadratic demonstration output")
    add_figure_placeholder(doc, "Figure 6.3 – Compare mode: iteration/call counts and ASCII contours")
    add_figure_placeholder(doc, "Figure 6.4 – Real-data training comparison (UCI)")
    add_figure_placeholder(doc, "Figure 6.5 – Held-out metrics and confusion counts")
    add_figure_placeholder(doc, "Figure 6.6 – Top-25 ranking output")
    add_figure_placeholder(doc, "Figure 6.7 – Cost/threshold change without retraining")
    add_figure_placeholder(doc, "Figure 6.8 – Sample verification of a rejected/invalid input case")
    doc.add_page_break()

    # ----- 7 Contribution -----
    add_h1(doc, "7 INDIVIDUAL CONTRIBUTION")
    add_para(
        doc,
        "Fill in the genuine division of work and insert a screenshot of the Git "
        "contribution / effort analysis (for example GitHub insights or a git shortlog "
        "summary). Do not fabricate commit history.",
    )
    add_para(doc, "Table 7.1 – Contribution summary (fill in)", bold=True, align="center", size=11)
    add_table(
        doc,
        ["Member", "Areas of work", "Estimated share"],
        [
            [STUDENT_A, "[theory / GD / report sections / …]", "[xx %]"],
            [STUDENT_B, "[Nesterov / data pipeline / tests / …]", "[xx %]"],
        ],
    )
    add_figure_placeholder(doc, "Figure 7.1 – Effort analysis (git / GitHub contribution evidence)")
    doc.add_page_break()

    # ----- 8 User guide -----
    add_h1(doc, "8 USER GUIDE")
    add_h2(doc, "8.1 About the program")
    add_para(
        doc,
        "RetentionRadar is a console tool that trains logistic regression with classical "
        "gradient descent and Nesterov acceleration, then evaluates and ranks held-out "
        "student records for an educational retention demonstration.",
    )
    add_h2(doc, "8.2 Requirements")
    add_bullets(
        doc,
        [
            "JDK 17 or newer on PATH (or selected in IntelliJ);",
            "working directory set to the project root so data/students-first-semester.csv resolves;",
            "no external Java libraries required.",
        ],
    )
    add_h2(doc, "8.3 Run in IntelliJ")
    add_para(
        doc,
        "Open the project folder, select JDK 17+, set the working directory to the project "
        "root, and run auk.sdt202.team4.app.RetentionRadarApp. With no program arguments "
        "the interactive menu appears. Optional arguments: --verify, --quadratic, --compare, "
        "--synthetic, --real, or --real \"path/to/file.csv\".",
    )
    add_h2(doc, "8.4 Run from a terminal")
    add_para(
        doc,
        "mkdir -p build/classes && find src -name '*.java' > sources.txt && "
        "javac --release 17 -d build/classes @sources.txt && "
        "java -Dfile.encoding=UTF-8 -cp build/classes auk.sdt202.team4.app.RetentionRadarApp",
        size=10,
        align="left",
    )
    add_para(doc, "On Windows, build.ps1 / run.ps1 wrap the same workflow.")

    add_h2(doc, "8.5 Menu overview")
    add_para(
        doc,
        "The menu starts on the reproducible synthetic cohort. Changing cohorts clears the "
        "latest fitted model.",
    )
    add_bullets(
        doc,
        [
            "1 — Load bundled UCI cohort;",
            "2 — Load synthetic cohort;",
            "3 — Train classical gradient descent;",
            "4 — Train Nesterov accelerated gradient;",
            "5 — Compare both methods and evaluate held-out records;",
            "6 — Show held-out ranking (top 25);",
            "7 — Change illustrative costs and threshold (no retraining);",
            "8 — Quadratic comparison and ASCII contours;",
            "9 — Run verification checks;",
            "0 — Exit.",
        ],
    )

    add_h2(doc, "8.6 Option 1 — Load UCI cohort")
    add_para(
        doc,
        "Loads data/students-first-semester.csv, omits Enrolled outcomes, and prepares a "
        "stratified split with train-only standardization when training is requested.",
    )
    add_figure_placeholder(doc, "Figure 8.1 – Loading the UCI cohort")

    add_h2(doc, "8.7 Option 5 — Compare and evaluate")
    add_para(
        doc,
        "Trains both optimizers with the same objective, initial point, step and tolerance, "
        "prints call counts and held-out metrics, and draws projected ASCII contours.",
    )
    add_figure_placeholder(doc, "Figure 8.2 – Comparison and evaluation output")

    add_h2(doc, "8.8 Option 7 — Decision costs")
    add_para(
        doc,
        "Enter positive costs for a missed positive outcome and for unnecessary outreach. "
        "The score threshold is derived from these illustrative costs; model weights do not "
        "change. Costs are demonstration assumptions, not evidence that outreach prevents dropout.",
    )
    add_figure_placeholder(doc, "Figure 8.3 – Changing costs and observing the new threshold")

    add_h2(doc, "8.9 Option 8 — Quadratic contours")
    add_para(
        doc,
        "Compares both methods on a known two-dimensional quadratic and prints ASCII "
        "contours with sampled ordinary iterates.",
    )
    add_figure_placeholder(doc, "Figure 8.4 – Quadratic ASCII contour view")

    add_h2(doc, "8.10 Option 9 — Verification")
    add_para(
        doc,
        "Runs optimizer, model, acceleration, data-pipeline and evaluation checks. All "
        "checks should report PASS before a demonstration or submission.",
    )
    add_figure_placeholder(doc, "Figure 8.5 – Verification menu result")
    doc.add_page_break()

    # ----- 9 Abstract -----
    add_h1(doc, "9 ABSTRACT")
    add_para(
        doc,
        "Gradient Descent Algorithm — RetentionRadar",
        align="center",
        bold=True,
        size=13,
        space_after=12,
    )
    add_para(
        doc,
        f"Authors: {STUDENT_A}, student of group {GROUP}; "
        f"{STUDENT_B}, student of group {GROUP}",
        align="center",
        size=11,
        space_after=6,
    )
    add_para(doc, "American University Kyiv, Kyiv, Ukraine", align="center", size=11, space_after=18)
    add_para(
        doc,
        "This project implements classical batch gradient descent and Nesterov accelerated "
        "gradient descent in plain Java 17 and applies both methods to logistic regression. "
        "The educational application RetentionRadar estimates the probability of eventual "
        "student dropout versus graduation from seven first-semester academic and "
        "administrative features, using a reproducible synthetic cohort and an attributed "
        "UCI data extract. The shared optimizer interface returns ordinary iterates with "
        "explicit stop reasons. On a fixed stratified UCI split, both methods reached the "
        "same held-out accuracy and AUC to the displayed precision, while Nesterov required "
        "fewer parameter updates (with additional look-ahead evaluations). The deliverable "
        "includes ASCII trajectory contours, held-out ranking, an illustrative cost-based "
        "threshold, and a 58-check offline verification suite. Limitations of the selected "
        "cohort, feature timing and single-split evaluation are stated explicitly.",
    )
    add_para(
        doc,
        f"Supervisor: {INSTRUCTOR}",
        align="right",
        size=11,
        space_after=18,
    )
    doc.add_page_break()

    # ----- References -----
    add_h1(doc, "REFERENCES")
    refs = [
        (
            "1. Tibshirani R. Gradient Descent. Carnegie Mellon University convex optimization "
            "lecture notes. URL: https://www.stat.cmu.edu/~ryantibs/convexopt/lectures/grad-descent.pdf "
            "(accessed: 09.10.2026)."
        ),
        (
            "2. Su W., Boyd S., Candes E. A Differential Equation for Modeling Nesterov’s "
            "Accelerated Gradient Method: Theory and Insights. URL: "
            "https://web.stanford.edu/~boyd/papers/pdf/ode_nest_grad.pdf (accessed: 09.10.2026)."
        ),
        (
            "3. Realinho V. et al. Predict Students’ Dropout and Academic Success. UCI Machine "
            "Learning Repository, dataset #697. URL: https://archive.ics.uci.edu/dataset/697/"
            "predict+students+dropout+and+academic+success (accessed: 09.10.2026)."
        ),
        (
            "4. Realinho V., Machado J., Baptista L., Martins M. V. Predicting Student Dropout "
            "and Academic Success. Data 2022, 7(11), 146. URL: https://www.mdpi.com/2306-5729/7/11/146 "
            "(accessed: 09.10.2026)."
        ),
        (
            "5. Shalizi C. Logistic Regression. CMU lecture notes. URL: "
            "https://stat.cmu.edu/~cshalizi/dm/20/lectures/07/lecture-07.html (accessed: 09.10.2026)."
        ),
        (
            "6. Creative Commons Attribution 4.0 International (CC BY 4.0). URL: "
            "https://creativecommons.org/licenses/by/4.0/ (accessed: 09.10.2026)."
        ),
        (
            "7. Project repository (GitHub). URL: "
            f"{REPO_URL} (accessed: 09.10.2026)."
        ),
    ]
    for r in refs:
        # Make URL portions clickable where practical for ref 3 and 7
        add_para(doc, r, align="justify", size=11, space_after=8)

    # Explicit clickable line for repository (brief requirement)
    p = doc.add_paragraph()
    style_paragraph(p, space_after=8)
    set_run_font(p.add_run("Clickable repository link: "), size=11)
    add_hyperlink(p, REPO_URL, REPO_URL)

    p = doc.add_paragraph()
    style_paragraph(p, space_after=8)
    set_run_font(p.add_run("Clickable UCI dataset link: "), size=11)
    add_hyperlink(
        p,
        "https://archive.ics.uci.edu/dataset/697/predict+students+dropout+and+academic+success",
        "https://archive.ics.uci.edu/dataset/697/predict+students+dropout+and+academic+success",
    )
    doc.add_page_break()

    # ----- Appendix A -----
    add_h1(doc, "APPENDIX A  SOURCE CODE LISTING")
    add_para(
        doc,
        f"{STUDENT_A}; {STUDENT_B}\n"
        "SDT 202 — Design and Analysis of Data Structures and Algorithms\n"
        "Project: Gradient Descent Algorithm — RetentionRadar\n"
        "Date: [DD Month YYYY]\n"
        "Estimated team hours: [N hours]\n"
        "We acknowledge that this is our team work.",
        size=11,
        align="left",
        space_after=12,
    )
    add_para(
        doc,
        "Insert the required program header and human-authored comments into the Java "
        "sources, then paste the listings below (or attach the repository tree). The file "
        "inventory of the current foundation is:",
    )
    files = [
        "auk/sdt202/team4/app/RetentionRadarApp.java — driver and interactive menu",
        "auk/sdt202/team4/app/ComparisonReport.java — comparison tables and contour projection",
        "auk/sdt202/team4/app/Evaluation.java — held-out metrics",
        "auk/sdt202/team4/app/Watchlist.java — ranking",
        "auk/sdt202/team4/app/DecisionCosts.java — illustrative threshold from costs",
        "auk/sdt202/team4/app/AsciiContour.java — ASCII contour renderer",
        "auk/sdt202/team4/app/QuadraticObjective.java — known quadratic demo objective",
        "auk/sdt202/team4/optimization/ObjectiveFunction.java",
        "auk/sdt202/team4/optimization/OptimizerConfig.java",
        "auk/sdt202/team4/optimization/OptimizationResult.java",
        "auk/sdt202/team4/optimization/StopReason.java",
        "auk/sdt202/team4/optimization/GradientDescent.java",
        "auk/sdt202/team4/optimization/NesterovGradientDescent.java",
        "auk/sdt202/team4/optimization/VectorMath.java",
        "auk/sdt202/team4/model/LogisticObjective.java",
        "auk/sdt202/team4/data/SyntheticCohort.java",
        "auk/sdt202/team4/data/StudentCsvLoader.java",
        "auk/sdt202/team4/data/DataSplit.java",
        "auk/sdt202/team4/data/Standardizer.java",
        "auk/sdt202/team4/verify/SelfCheck.java",
        "auk/sdt202/team4/verify/ModelCheck.java",
        "auk/sdt202/team4/verify/AccelerationCheck.java",
        "auk/sdt202/team4/verify/DataCheck.java",
        "auk/sdt202/team4/verify/EvaluationCheck.java",
    ]
    add_bullets(doc, files)

    add_para(
        doc,
        "Representative classical update (conceptual; match GradientDescent.java after "
        "adding your comments):",
        space_after=6,
    )
    add_para(
        doc,
        "x ← x − α · ∇f(x)\n"
        "stop when ‖∇f(x)‖ ≤ tolerance, or after maxUpdates, or on non-finite / magnitude divergence",
        size=10,
        italic=True,
        align="left",
    )
    add_para(
        doc,
        "Representative Nesterov look-ahead step (conceptual; match NesterovGradientDescent.java):",
    )
    add_para(
        doc,
        "evaluate gradient at look-ahead y;\n"
        "x ← y − α · ∇f(y);\n"
        "β ← (k−1)/(k+2);\n"
        "y ← x + β (x − x_previous);\n"
        "return and stop on the ordinary iterate x",
        size=10,
        italic=True,
        align="left",
    )
    add_para(
        doc,
        "[PASTE FULL JAVA SOURCE LISTINGS HERE — one subsection per file, after human "
        "comments/header are written.]",
        italic=True,
        align="center",
        size=11,
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
